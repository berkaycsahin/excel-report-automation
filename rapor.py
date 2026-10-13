"""Şubelerden gelen dağınık Excel dosyalarını birleştirip tek bir özet rapor üretir.

Kullanım:
    python rapor.py                      # girdi/ klasöründeki tüm .xlsx dosyaları
    python rapor.py --girdi D:/satislar --cikti rapor.xlsx
"""
import argparse
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Şubeler aynı bilgiyi farklı sütun adlarıyla gönderiyor; hepsini tek ada eşliyoruz
SUTUN_ESLESME = {
    "tarih": "Tarih",
    "ürün": "Ürün", "urun": "Ürün", "urun_adi": "Ürün", "ürün adı": "Ürün",
    "adet": "Adet", "miktar": "Adet",
    "birim fiyat": "Fiyat", "fiyat": "Fiyat", "fiyat (tl)": "Fiyat",
}

BASLIK_DOLGU = PatternFill("solid", fgColor="1F4E78")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")
KENAR = Border(bottom=Side(style="thin", color="BFBFBF"))
TL = '#,##0.00 "₺"'


def sutun_adi(ad: str) -> str:
    anahtar = str(ad).strip().replace("İ", "i").replace("I", "ı").lower()
    return SUTUN_ESLESME.get(anahtar, str(ad).strip())


def tarih_coz(deger):
    """2026-09-01 (ISO) yıl başta olduğu için ayrı okunur; 01.09.2026 ve 01/09/2026 gün başta kabul edilir."""
    if isinstance(deger, str) and len(deger.strip()) >= 8 and deger.strip()[4] == "-":
        return pd.to_datetime(deger.strip(), format="%Y-%m-%d", errors="coerce")
    return pd.to_datetime(deger, dayfirst=True, errors="coerce")


def dosya_oku(yol: Path) -> pd.DataFrame:
    df = pd.read_excel(yol)
    df.columns = [sutun_adi(c) for c in df.columns]
    eksik = {"Tarih", "Ürün", "Adet", "Fiyat"} - set(df.columns)
    if eksik:
        raise ValueError(f"{yol.name}: tanınmayan sütunlar, eksik: {', '.join(sorted(eksik))}")
    df = df[["Tarih", "Ürün", "Adet", "Fiyat"]].copy()
    df["Şube"] = yol.stem.split("_")[0]
    return df


def temizle(df: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    log = {"Okunan satır": len(df)}

    df = df.dropna(how="all", subset=["Tarih", "Ürün", "Adet", "Fiyat"])
    log["Boş satır silindi"] = log["Okunan satır"] - len(df)

    df["Tarih"] = df["Tarih"].map(tarih_coz)
    df["Adet"] = pd.to_numeric(df["Adet"], errors="coerce")
    df["Fiyat"] = pd.to_numeric(df["Fiyat"], errors="coerce")
    hatali = df[["Tarih", "Adet", "Fiyat"]].isna().any(axis=1)
    log["Hatalı satır atlandı"] = int(hatali.sum())
    df = df[~hatali]

    # "çay 1kg", " Çay 1kg " ve "Çay 1kg" aynı ürün: en sık görülen yazımı kullan
    df["Ürün"] = df["Ürün"].astype(str).str.split().str.join(" ")
    anahtar = df["Ürün"].str.casefold()
    standart = df.groupby(anahtar)["Ürün"].agg(lambda s: s.mode().iloc[0])
    df["Ürün"] = anahtar.map(standart)

    once = len(df)
    df = df.drop_duplicates()
    log["Tekrarlanan kayıt silindi"] = once - len(df)

    df["Tutar"] = df["Adet"] * df["Fiyat"]
    log["Rapora giren satır"] = len(df)
    return df.sort_values(["Tarih", "Şube"]).reset_index(drop=True), log


def sayfa_bicimle(ws, para_sutunlari=(), genislik=None):
    for hucre in ws[1]:
        hucre.fill, hucre.font = BASLIK_DOLGU, BASLIK_YAZI
        hucre.alignment = Alignment(horizontal="center")
    for satir in ws.iter_rows(min_row=2):
        for hucre in satir:
            hucre.border = KENAR
            if hucre.column_letter in para_sutunlari:
                hucre.number_format = TL
    for i, sutun in enumerate(ws.columns, start=1):
        en_uzun = max(len(str(h.value or "")) for h in sutun)
        ws.column_dimensions[get_column_letter(i)].width = genislik or min(en_uzun + 4, 40)
    ws.freeze_panes = "A2"


def rapor_yaz(df: pd.DataFrame, log: dict, cikti: Path):
    sube = df.groupby("Şube").agg(Satış_Adedi=("Adet", "sum"), Ciro=("Tutar", "sum")).reset_index()
    sube = sube.sort_values("Ciro", ascending=False).rename(columns={"Satış_Adedi": "Satış Adedi"})
    urun = df.groupby("Ürün").agg(Adet=("Adet", "sum"), Ciro=("Tutar", "sum")).reset_index()
    urun = urun.sort_values("Ciro", ascending=False)
    gunluk = df.groupby(df["Tarih"].dt.date)["Tutar"].sum().reset_index()
    gunluk.columns = ["Tarih", "Ciro"]
    pivot = df.pivot_table(index="Ürün", columns="Şube", values="Tutar", aggfunc="sum", fill_value=0)

    with pd.ExcelWriter(cikti, engine="openpyxl") as yazici:
        sube.to_excel(yazici, sheet_name="Özet", index=False, startrow=6)
        urun.to_excel(yazici, sheet_name="Ürünler", index=False)
        pivot.reset_index().to_excel(yazici, sheet_name="Ürün x Şube", index=False)
        gunluk.to_excel(yazici, sheet_name="Günlük", index=False)
        ham = df.assign(Tarih=df["Tarih"].dt.date)
        ham.to_excel(yazici, sheet_name="Temiz Veri", index=False)
        pd.DataFrame(list(log.items()), columns=["Adım", "Satır"]).to_excel(
            yazici, sheet_name="İşlem Kaydı", index=False)

        wb = yazici.book
        ozet = wb["Özet"]
        ozet["A1"] = "Satış Raporu"
        ozet["A1"].font = Font(bold=True, size=16, color="1F4E78")
        ozet["A2"] = (f"{df['Tarih'].min():%d.%m.%Y} – {df['Tarih'].max():%d.%m.%Y}  ·  "
                      f"{df['Şube'].nunique()} şube  ·  oluşturma: {date.today():%d.%m.%Y}")
        ozet["A2"].font = Font(italic=True, color="7F7F7F")
        for i, (etiket, deger, bicim) in enumerate([
            ("Toplam Ciro", df["Tutar"].sum(), TL),
            ("Toplam Satış Adedi", df["Adet"].sum(), "#,##0"),
            ("En Çok Satan Ürün", urun.iloc[0]["Ürün"], None),
        ]):
            ozet.cell(row=4, column=1 + i * 2, value=etiket).font = Font(color="7F7F7F")
            h = ozet.cell(row=5, column=1 + i * 2, value=deger)
            h.font = Font(bold=True, size=13)
            if bicim:
                h.number_format = bicim
        for hucre in ozet[7]:
            hucre.fill, hucre.font = BASLIK_DOLGU, BASLIK_YAZI
        for satir in ozet.iter_rows(min_row=8, min_col=3, max_col=3):
            satir[0].number_format = TL
        for harf in "ABCDEF":
            ozet.column_dimensions[harf].width = 20

        grafik = BarChart()
        grafik.title, grafik.y_axis.title = "Şubelere Göre Ciro", "₺"
        grafik.add_data(Reference(ozet, min_col=3, min_row=7, max_row=7 + len(sube)), titles_from_data=True)
        grafik.set_categories(Reference(ozet, min_col=1, min_row=8, max_row=7 + len(sube)))
        grafik.legend, grafik.height, grafik.width = None, 7, 14
        ozet.add_chart(grafik, f"A{10 + len(sube)}")

        sayfa_bicimle(wb["Ürünler"], para_sutunlari=("C",))
        uw = wb["Ürünler"]
        g = BarChart()
        g.type, g.title, g.legend = "bar", "Ürünlere Göre Ciro", None
        g.add_data(Reference(uw, min_col=3, min_row=1, max_row=len(urun) + 1), titles_from_data=True)
        g.set_categories(Reference(uw, min_col=1, min_row=2, max_row=len(urun) + 1))
        g.height, g.width = 8, 14
        uw.add_chart(g, "E2")

        pw = wb["Ürün x Şube"]
        sayfa_bicimle(pw, para_sutunlari=tuple(get_column_letter(i) for i in range(2, pw.max_column + 1)))

        gw = wb["Günlük"]
        sayfa_bicimle(gw, para_sutunlari=("B",))
        for (h,) in gw.iter_rows(min_row=2, max_col=1):
            h.number_format = "dd.mm.yyyy"
        cizgi = LineChart()
        cizgi.title, cizgi.legend, cizgi.height, cizgi.width = "Günlük Ciro", None, 8, 18
        cizgi.add_data(Reference(gw, min_col=2, min_row=1, max_row=len(gunluk) + 1), titles_from_data=True)
        cizgi.set_categories(Reference(gw, min_col=1, min_row=2, max_row=len(gunluk) + 1))
        cizgi.x_axis.number_format = "dd.mm"
        gw.add_chart(cizgi, "D2")

        tw = wb["Temiz Veri"]
        sayfa_bicimle(tw, para_sutunlari=("D", "F"))
        for (h,) in tw.iter_rows(min_row=2, max_col=1):
            h.number_format = "dd.mm.yyyy"
        tw.auto_filter.ref = tw.dimensions
        sayfa_bicimle(wb["İşlem Kaydı"])


def main():
    p = argparse.ArgumentParser(description="Şube Excel dosyalarını birleştirip rapor üretir.")
    p.add_argument("--girdi", type=Path, default=Path(__file__).parent / "girdi")
    p.add_argument("--cikti", type=Path,
                   default=Path(__file__).parent / "cikti" / f"rapor_{date.today():%Y-%m-%d}.xlsx")
    a = p.parse_args()

    dosyalar = sorted(f for f in a.girdi.glob("*.xlsx") if not f.name.startswith("~$"))
    if not dosyalar:
        raise SystemExit(f"'{a.girdi}' içinde .xlsx dosyası bulunamadı.")

    df, log = temizle(pd.concat([dosya_oku(f) for f in dosyalar], ignore_index=True))
    a.cikti.parent.mkdir(parents=True, exist_ok=True)
    rapor_yaz(df, log, a.cikti)

    print(f"{len(dosyalar)} dosya işlendi:")
    for adim, sayi in log.items():
        print(f"  {adim:<28}{sayi:>6}")
    print(f"Rapor: {a.cikti}")


if __name__ == "__main__":
    main()
